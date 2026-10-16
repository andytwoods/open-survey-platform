"""Write XLSForm sheets to .xlsx bytes with openpyxl, deterministically.

openpyxl stamps the current time into the document properties and the zip entries; both
are pinned here, so the same survey always gives the same bytes, whatever the clock says.
"""

from __future__ import annotations

import datetime as _dt
import io
import zipfile
from collections.abc import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.writer.excel import ExcelWriter

_FIXED = _dt.datetime(2000, 1, 1)
_ZIP_TIME = (2000, 1, 1, 0, 0, 0)

type Sheet = Sequence[Mapping[str, str]]


def _repack(data: bytes) -> bytes:
    """Rewrite the zip with fixed entry timestamps and permissions."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            dst.writestr(info, src.read(item.filename))
    return out.getvalue()


def columns(rows: Sheet, preferred: Sequence[str]) -> list[str]:
    """Columns used by ``rows``: ``preferred`` order first, then first-seen order."""
    used: dict[str, None] = {}
    for row in rows:
        for k, v in row.items():
            if v != "":
                used.setdefault(k, None)
    head = [c for c in preferred if c in used]
    return head + [c for c in used if c not in head]


def write_workbook(sheets: Mapping[str, tuple[list[str], Sheet]]) -> bytes:
    """``sheets`` maps a sheet name to ``(header, rows)``, in sheet order."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        for r, values in enumerate([dict(zip(header, header, strict=True)), *rows], start=1):
            for c, col in enumerate(header, start=1):
                # Control characters cannot be stored in .xlsx; the exporter reports them.
                value = ILLEGAL_CHARACTERS_RE.sub("", values.get(col, ""))
                if value:
                    cell = ws.cell(row=r, column=c, value=value)
                    cell.data_type = "s"  # never a formula, even for text starting with "="
    wb.properties.created = _FIXED
    wb.properties.modified = _FIXED
    wb.properties.creator = "osp-targets"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as archive:
        ExcelWriter(wb, archive).write_data()
    return _repack(buf.getvalue())
