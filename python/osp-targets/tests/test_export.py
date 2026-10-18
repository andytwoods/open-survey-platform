"""XLSForm export: every survey converts with pyxform, and every non-native feature used
is reported with a JSON pointer that resolves in the survey document.

Surveys: every fixture in /fixtures (osp-survey/0.1 since C3b) plus the 10 valid cases in
python/osp-spec/tests/cases/valid.
"""

from __future__ import annotations

import copy
import io
import json
import os
import zipfile
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from osp_spec import parse
from osp_targets import XLSFORM, build_xlsform, export_xlsform
from osp_targets.xlsform import feature_uses
from pyxform.xls2xform import convert

REPO = Path(__file__).resolve().parents[3]
CASES = sorted((REPO / "python/osp-spec/tests/cases/valid").glob("*.json"))
FIXTURES = sorted((REPO / "fixtures").glob("*/survey.json"))
GOLDEN = Path(__file__).parent / "data/xlsform"
GOLDEN_FIXTURES = ["27-unicode-text", "41-rich-text"]


def _fixture_survey(path: Path) -> Any:
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("survey", data) if isinstance(data, dict) else data


def _loads(data: Any) -> bool:
    survey, diags = parse(data)
    return survey is not None and not any(d.severity == "error" for d in diags)


SURVEYS: list[tuple[str, Any]] = [
    (f"case/{p.stem}", json.loads(p.read_text(encoding="utf-8"))["survey"]) for p in CASES
] + [(f"fixture/{p.parent.name}", _fixture_survey(p)) for p in FIXTURES]


def _xform(data: bytes) -> str:
    warnings: list[str] = []
    result = convert(xlsform=io.BytesIO(data), warnings=warnings, file_type=".xlsx")
    return result.xform


def _resolve(doc: Any, ptr: str) -> Any:
    if ptr == "":
        return doc
    for raw in ptr.split("/")[1:]:
        part = raw.replace("~1", "/").replace("~0", "~")
        doc = doc[int(part)] if isinstance(doc, list) else doc[part]
    return doc


def test_there_are_surveys() -> None:
    assert len(CASES) >= 10


def test_every_fixture_loads() -> None:
    assert len(FIXTURES) >= 20
    bad = [p.parent.name for p in FIXTURES if not _loads(_fixture_survey(p))]
    assert bad == []


@pytest.mark.parametrize(("label", "data"), SURVEYS, ids=[s[0] for s in SURVEYS])
def test_export_converts(label: str, data: Any) -> None:
    xlsx, diags = export_xlsform(data)
    assert not [d for d in diags if d.severity == "error"], diags
    xform = _xform(xlsx)
    assert '<h:body class="pages">' in xform
    assert 'appearance="field-list"' in xform

    doc = json.loads(json.dumps(data))
    for d in diags:
        assert d.target == "xlsform"
        assert d.feature in XLSFORM, d
        cap = XLSFORM[d.feature]
        assert d.level == cap.level != "native"
        assert d.code == f"xlsform.{cap.level}"
        assert d.severity == {"lowered": "info", "approximated": "warning",
                              "unsupported": "loss"}[cap.level]  # fmt: skip
        _resolve(doc, d.pointer)  # raises if the pointer does not resolve

    # Every non-native field the survey uses has a diagnostic at its pointer.
    survey, _ = parse(data)
    reported = {(d.feature, d.pointer) for d in diags}
    for key, ptr, _value in feature_uses(survey):
        if XLSFORM[key].level != "native":
            assert (key, ptr) in reported, (key, ptr)


@pytest.mark.parametrize("name", GOLDEN_FIXTURES)
def test_golden(name: str) -> None:
    """The sheets and diagnostics of a fixture, pinned (regenerate with
    ``OSP_UPDATE_GOLDEN=1``): 27 for plain text with literal markup, 41 for rich text."""
    form, diags = build_xlsform(_fixture_survey(REPO / "fixtures" / name / "survey.json"))
    assert form is not None, diags
    sheets = {s: {"columns": cols, "rows": rows} for s, (cols, rows) in form.sheets().items()}
    got_sheets = json.dumps(sheets, indent=1, ensure_ascii=False) + "\n"
    got_diags = "".join(f"{d.severity}\t{d.feature}\t{d.pointer}\n" for d in diags)
    path_sheets = GOLDEN / f"{name}.sheets.json"
    path_diags = GOLDEN / f"{name}.diagnostics.tsv"
    if os.environ.get("OSP_UPDATE_GOLDEN") == "1":
        GOLDEN.mkdir(parents=True, exist_ok=True)
        path_sheets.write_text(got_sheets, encoding="utf-8", newline="\n")
        path_diags.write_text(got_diags, encoding="utf-8", newline="\n")
    assert got_sheets == path_sheets.read_text(encoding="utf-8"), "OSP_UPDATE_GOLDEN=1"
    assert got_diags == path_diags.read_text(encoding="utf-8")


def test_deterministic_bytes() -> None:
    data = json.loads(CASES[0].read_text(encoding="utf-8"))["survey"]
    a, _ = export_xlsform(data)
    b, _ = export_xlsform(copy.deepcopy(data))
    assert a == b
    with zipfile.ZipFile(io.BytesIO(a)) as z:
        assert {i.date_time for i in z.infolist()} == {(2000, 1, 1, 0, 0, 0)}


def test_accepts_survey_model() -> None:
    data = json.loads(CASES[0].read_text(encoding="utf-8"))["survey"]
    survey, _ = parse(data)
    assert export_xlsform(survey) == export_xlsform(data)


# --- Feature checks on the worked example and small surveys -----------------------------


def _survey(items: list[dict], flow_extra: list[dict] | None = None, **top: Any) -> dict:
    return {
        "specVersion": "osp-survey/0.1",
        "id": "s1",
        "metadata": {"title": "T", "languages": ["en"]},
        "blocks": [{"id": "b1", "pages": [{"id": "p1", "items": items}]}],
        "flow": [{"type": "block", "id": "f1", "block": "b1"}, *(flow_extra or [])],
        **top,
    }


FRUIT = [{"value": "apple", "label": "Apple"}, {"value": "pear", "label": "Pear"}]


def test_piped_choice_label_uses_choice_name() -> None:
    data = _survey(
        [
            {"id": "fav", "type": "single", "label": "Fruit?", "choices": FRUIT},
            {"id": "n", "type": "note", "label": "You chose ${fav}."},
        ]
    )
    xlsx, diags = export_xlsform(data)
    xform = _xform(xlsx)
    assert "jr:choice-name( /data/b1/p1/fav , ' /data/b1/p1/fav ')" in xform
    assert '<output value=" /data/b1/p1/osp_label_fav "/>' in xform
    assert ("$pipe.choiceLabel", "/blocks/0/pages/0/items/1/label") in {
        (d.feature, d.pointer) for d in diags
    }
    assert ("SingleQuestion.choices", "/blocks/0/pages/0/items/0/choices") in {
        (d.feature, d.pointer) for d in diags
    }


def test_choice_shuffle_is_seeded_randomize() -> None:
    data = _survey(
        [
            {
                "id": "fav",
                "type": "multi",
                "label": "Fruit?",
                "choices": FRUIT,
                "choiceOrder": "shuffle",
            }
        ]  # fmt: skip
    )
    xlsx, diags = export_xlsform(data)
    xform = _xform(xlsx)
    assert "randomize(instance('fav_choices')/root/item, /data/b1/p1/osp_seed_fav)" in xform
    assert "once(int(random() * 2147483646) + 1)" in xform
    assert ("MultiQuestion.choiceOrder=shuffle", "/blocks/0/pages/0/items/0/choiceOrder") in {
        (d.feature, d.pointer) for d in diags
    }


def test_end_node_and_outcome() -> None:
    data = _survey(
        [{"id": "ok", "type": "single", "label": "OK?", "choices": FRUIT}],
        [
            {
                "type": "branch",
                "id": "g",
                "if": {
                    "kind": "compare",
                    "op": "==",
                    "left": {"kind": "ref", "id": "ok"},
                    "right": {"kind": "text", "value": "pear"},
                },  # fmt: skip
                "then": [{"type": "end", "id": "e1", "outcome": "screened", "message": "Bye"}],
            }
        ],
        blocks=[
            {
                "id": "b1",
                "pages": [
                    {
                        "id": "p1",
                        "items": [{"id": "ok", "type": "single", "label": "OK?", "choices": FRUIT}],
                    }
                ],
            },
            {
                "id": "b2",
                "pages": [{"id": "p2", "items": [{"id": "more", "type": "text", "label": "More"}]}],
            },
        ],  # fmt: skip
    )
    data["flow"].append({"type": "block", "id": "f2", "block": "b2"})
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    assert rows["b2"]["relevant"].startswith("not(")
    assert "'screened'" in rows["osp_outcome"]["calculation"]
    assert rows["osp_end_e1"]["relevant"] == rows["b2"]["relevant"].removeprefix("not(")[:-1]
    xlsx, _ = export_xlsform(data)
    _xform(xlsx)
    assert "EndNode.type" in {d.feature for d in diags}


def test_endings_become_note_groups() -> None:
    """EN1: another outcome's ending goes with its End node (title = group label, body after
    the node's message); the `complete` ending is the final group."""
    data = _survey(
        [{"id": "ok", "type": "single", "label": "OK?", "choices": FRUIT}],
        [
            {
                "type": "branch",
                "id": "g",
                "if": {
                    "kind": "compare",
                    "op": "==",
                    "left": {"kind": "ref", "id": "ok"},
                    "right": {"kind": "text", "value": "pear"},
                },  # fmt: skip
                "then": [{"type": "end", "id": "e1", "outcome": "screened", "message": "Bye"}],
            }
        ],
        settings={
            "endings": {
                "complete": {"title": "Thanks", "body": "You chose ${ok}."},
                "screened": {"title": "Sorry", "body": "Not this time.", "status": "none"},
            }
        },
    )
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    names = [r.get("name") for r in form.survey]
    assert rows["osp_end_e1"]["label::English (en)"] == "Sorry"
    assert names.index("osp_end_e1_message") < names.index("osp_end_e1_ending")
    assert rows["osp_end_e1_ending"]["label::English (en)"] == "Not this time."
    final = rows["osp_ending_complete"]
    assert final["label::English (en)"] == "Thanks"
    assert final["relevant"] == rows["osp_end_e1"]["relevant"].join(["not(", ")"])
    # A piped choice shows its label, as in question labels (semantics.md §8).
    assert rows["osp_ending_complete_body"]["label::English (en)"] == "You chose ${osp_label_ok}."
    levels = {d.feature: d.level for d in diags}
    assert levels["Settings.endings"] == "lowered" and levels["Ending.status"] == "approximated"
    xlsx, _ = export_xlsform(data)
    _xform(xlsx)


def test_unsupported_features_are_losses() -> None:
    data = _survey(
        [{"id": "a", "type": "integer", "label": "A"}],
        settings={"backNavigation": False},
    )
    data["metadata"]["description"] = "D"
    _, diags = export_xlsform(data)
    losses = {(d.feature, d.pointer) for d in diags if d.severity == "loss"}
    assert ("Settings.backNavigation=false", "/settings/backNavigation") in losses
    assert ("Metadata.description", "/metadata/description") in losses


def test_missing_states_documented() -> None:
    _, diags = export_xlsform(_survey([{"id": "a", "type": "integer", "label": "A"}]))
    states = [d for d in diags if d.feature == "$missingStates"]
    assert len(states) == 1 and states[0].pointer == ""
    for word in ("not-reached", "irrelevant", "unanswered"):
        assert word in states[0].message


def test_text_is_never_a_formula_and_control_chars_reported() -> None:
    data = _survey([{"id": "a", "type": "text", "label": "=1+1 \x07 *bold* <b>"}])
    xlsx, diags = export_xlsform(data)
    ws = load_workbook(io.BytesIO(xlsx))["survey"]
    labels = [c.value for row in ws.iter_rows() for c in row]
    # Literal markup is escaped (RT4, lead decision 5); `<` has no escape and is reported.
    assert "=1+1  \\*bold\\* <b>" in labels
    features = {d.feature for d in diags}
    assert {"$text.controlCharacters", "$text.markup"} <= features
    _xform(xlsx)


def _label(name: str, lang: str = "English (en)") -> Any:
    def get(xlsx: bytes, col: str = "label") -> str:
        ws = load_workbook(io.BytesIO(xlsx))["survey"]
        rows = list(ws.iter_rows(values_only=True))
        head = list(rows[0])
        row = next(r for r in rows[1:] if r[head.index("name")] == name)
        return row[head.index(f"{col}::{lang}")]

    return get


def test_rich_text_is_markdown_with_pipes_and_one_diagnostic_per_node() -> None:
    """RT4: marks, links, lists and pipes inside formatting (fixture 41 pins the rest)."""
    label = {
        "type": "rich",
        "blocks": [
            {
                "type": "p",
                "content": [
                    {"text": "Hello ", "marks": ["b"]},
                    {"pipe": "fav"},
                    {"text": ", "},
                    {"text": "Merci", "lang": "fr"},
                    {"text": " * ", "marks": ["i"]},
                    {"text": "sheet", "href": "https://example.org/"},
                ],
            },
            {"type": "ol", "items": [[{"text": "a"}, {"br": True}, {"text": "b"}]]},
        ],
    }
    fav = {"id": "fav", "type": "single", "label": "Fruit", "choices": FRUIT}
    data = _survey([fav, {"id": "n", "type": "note", "label": label}])
    xlsx, diags = export_xlsform(data)
    assert _label("n")(xlsx) == (
        "**Hello** ${osp_label_fav}, Merci *\\** [sheet](https://example.org/)\n\n1. a b"
    )
    base = "/blocks/0/pages/0/items/1/label"
    by = {(d.feature, d.pointer) for d in diags}
    assert ("TextRun.lang", base + "/blocks/0/content/3/lang") in by
    assert ("$pipe.choiceLabel", base + "/blocks/0/content/1") in by
    assert ("RichList.type", base + "/blocks/1") in by
    assert ("$rich.listBreak", base + "/blocks/1/items/0/1") in by
    assert not [d for d in diags if d.feature.startswith("$text.")]
    assert len(diags) == len({(d.feature, d.pointer) for d in diags})  # once per node
    _xform(xlsx)


def test_plain_text_markup_is_escaped_in_every_text_column() -> None:
    """Lead decision 5: literal `* _ #` in existing plain text show as written in
    Enketo/Collect; `$text.markup` is only for what has no escape."""
    data = _survey(
        [
            {
                "id": "f",
                "type": "single",
                "label": "F",
                "choices": [{"value": "v", "label": "a*b"}],
            },
            {
                "id": "a",
                "type": "integer",
                "label": "Score #1 of ${f}_x",
                "help": "_hint_",
                "constraint": {
                    "kind": "compare",
                    "op": "<=",
                    "left": {"kind": "ref", "id": "a"},
                    "right": {"kind": "int", "value": 5},
                },
                "constraintMessage": "5 * 1 max",
            },
        ]
    )
    xlsx, diags = export_xlsform(data)
    assert xlsx, diags
    get = _label("a")
    assert get(xlsx) == "Score \\#1 of ${osp_label_f}\\_x"
    assert get(xlsx, "hint") == "\\_hint\\_"
    assert get(xlsx, "constraint_message") == "5 \\* 1 max"
    ws = load_workbook(io.BytesIO(xlsx))["choices"]
    assert "a\\*b" in [c.value for row in ws.iter_rows() for c in row]
    assert "$text.markup" not in {d.feature for d in diags}
    _xform(xlsx)


def test_names_used_and_collisions_avoided() -> None:
    data = _survey(
        [
            {"id": "q1", "type": "integer", "label": "A", "name": "age"},
            {
                "id": "q2",
                "type": "single",
                "label": "B",
                "name": "osp_seed_q2",
                "choices": FRUIT,
                "choiceOrder": "shuffle",
            },  # fmt: skip
        ]
    )
    data["blocks"][0]["id"] = "b1"
    form, _ = build_xlsform(data)
    assert form is not None
    names = [r.get("name") for r in form.survey if r.get("name")]
    assert len(names) == len(set(names))
    assert "age" in names and "osp_seed_osp_seed_q2" in names


@pytest.mark.parametrize(
    "bad",
    [None, [], "x", 1, {}, {"specVersion": "osp-survey/0.1"}, {"blocks": [1, 2]}],
)
def test_invalid_input_gives_error_not_exception(bad: Any) -> None:
    data, diags = export_xlsform(bad)
    assert data == b""
    assert diags and all(d.severity == "error" for d in diags)


def test_mutated_cases_never_raise() -> None:
    """Hostile input: drop or replace every key in every valid case, one at a time."""
    for path in CASES:
        base = json.loads(path.read_text(encoding="utf-8"))["survey"]
        for key in list(base):
            for replacement in (None, [], {}, "x", -1):
                data = copy.deepcopy(base)
                data[key] = replacement
                xlsx, diags = export_xlsform(data)
                assert isinstance(xlsx, bytes) and isinstance(diags, list)
            data = copy.deepcopy(base)
            del data[key]
            export_xlsform(data)


def test_page_showif_min_max_and_if_are_lowered() -> None:
    """Decisions batch 1 (C2-9, C1-3): page relevant, min/max joined into constraint, if()."""
    gt = {"kind": "compare", "op": ">", "left": {"kind": "ref", "id": "n"},
          "right": {"kind": "int", "value": 1}}  # fmt: skip
    data = {
        "specVersion": "osp-survey/0.1",
        "id": "srv",
        "metadata": {"title": "T", "languages": ["en"]},
        "variables": [{"kind": "computed", "id": "band", "expr": {
            "kind": "if", "cond": gt, "then": {"kind": "text", "value": "hi"},
            "else": {"kind": "text", "value": "lo"}}}],  # fmt: skip
        "blocks": [
            {"id": "b1", "pages": [
                {"id": "p1", "items": [{"id": "n", "type": "integer", "label": "N",
                                        "min": 1, "max": 9, "constraint": gt}]},
                {"id": "p2", "showIf": gt, "items": [{"id": "t", "type": "text", "label": "T"}]},
            ]},
        ],
        "flow": [{"type": "block", "id": "f1", "block": "b1"}],
    }  # fmt: skip
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    assert rows["n"]["constraint"].startswith("(. >= 1) and (. <= 9) and (")
    assert "${n}" in rows["p2"]["relevant"]
    assert rows["band"]["calculation"].startswith("if(")
    content, _ = export_xlsform(data)
    assert "if(" in _xform(content)


def test_multiline_text_gets_multiline_appearance() -> None:
    data = _survey(
        [
            {"id": "long", "type": "text", "label": "Tell us", "multiline": True},
            {"id": "short", "type": "text", "label": "Name"},
        ]
    )
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    assert rows["long"]["appearance"] == "multiline"
    assert "appearance" not in rows["short"]
    content, _ = export_xlsform(data)
    assert 'appearance="multiline"' in _xform(content)


def test_matrix_is_a_table_list_and_slider_a_range() -> None:
    """C2c: matrix rows are select_one fields named <matrix name>_<row id> in a table-list
    group; a slider is a `range`; conditions on a row use the row's name."""
    data = json.loads((REPO / "fixtures/38-matrix/survey.json").read_text(encoding="utf-8"))
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    assert rows["drink"]["appearance"] == "table-list"
    assert rows["osp_drink_header"]["appearance"] == "label"
    assert rows["drink_tea"]["type"] == "select_one drink_choices"
    assert rows["drink_tea"].get("required") == "yes"
    assert "required" not in rows["drink_cocoa"]
    assert "${drink_tea}" in rows["cups"]["relevant"]
    assert rows["when"]["relevant"] == "string-length(${drink_cocoa}) > 0"
    assert rows["when_evening"].get("required") == "yes"

    data = json.loads((REPO / "fixtures/39-slider/survey.json").read_text(encoding="utf-8"))
    form, diags = build_xlsform(data)
    assert form is not None, diags
    rows = {r.get("name"): r for r in form.survey}
    assert rows["sleep"]["type"] == "range"
    assert rows["sleep"]["parameters"] == "start=0 end=12 step=0.5"
    assert "constraint" not in rows["pain"]  # range enforces min, max and step
    assert {d.feature for d in diags} >= {"SliderQuestion.minLabel", "SliderQuestion.maxLabel"}
